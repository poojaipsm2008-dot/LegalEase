"""Focused API, Gemini adapter, and export tests."""

import os
import sys
import types
import unittest
from io import BytesIO
from unittest.mock import patch
from zipfile import ZipFile

from fastapi.testclient import TestClient
from docx import Document
from pypdf import PdfReader

from ai_core.gemini_generator import GeminiDocumentGenerator, GeminiGenerationError
from main import app
import routes
from utils.document_formatter import format_docx
from utils.pdf_generator import format_pdf


class ApplicationTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    def test_health(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["status"], "ok")

    def test_validation_rejects_missing_and_blank_fields(self):
        self.assertEqual(self.client.post("/generate", json={}).status_code, 422)
        payload = {"document_type": "  ", "parties": "A, B", "terms": "Pay", "dates": "Today"}
        self.assertEqual(self.client.post("/generate", json=payload).status_code, 422)

    def test_swagger_and_generate_route(self):
        self.assertEqual(self.client.get("/docs").status_code, 200)

        class Generator:
            def generate_document(self, **kwargs):
                self.request = kwargs
                return "# NDA\n## Confidentiality\nBoth parties will protect information."

        with patch.object(routes, "GeminiDocumentGenerator", Generator):
            response = self.client.post("/generate", json={
                "document_type": "NDA", "parties": "Company and Consultant",
                "terms": "Protect confidential information", "dates": "10 April 2026",
            })
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["document_type"], "NDA")
        self.assertIn("Confidentiality", response.json()["content"])

    def test_generation_error_uses_safe_provider_status(self):
        class Generator:
            def __init__(self):
                pass

            def generate_document(self, **kwargs):
                raise GeminiGenerationError("Gemini is rate-limiting requests. Please wait and try again.", 429)

        with patch.object(routes, "GeminiDocumentGenerator", Generator):
            response = self.client.post("/generate", json={
                "document_type": "NDA", "parties": "A and B", "terms": "Keep data private", "dates": "Today",
            })
        self.assertEqual(response.status_code, 429)
        self.assertNotIn("api_key", response.text.lower())

    def test_missing_key_returns_service_unavailable_without_secret(self):
        with patch.dict(os.environ, {"GEMINI_API_KEY": ""}):
            response = self.client.post("/generate", json={
                "document_type": "NDA", "parties": "A and B", "terms": "Keep secrets", "dates": "Today"
            })
        self.assertEqual(response.status_code, 503)
        self.assertNotIn("unit-test-key", response.text)

    def test_gemini_uses_current_sdk_and_returns_text(self):
        calls = {}

        class Models:
            def generate_content(self, **kwargs):
                calls.update(kwargs)
                return types.SimpleNamespace(text="# NDA\n## Confidentiality\nPreserve this term.")

        class Client:
            def __init__(self, api_key):
                self.models = Models()

        fake_google = types.ModuleType("google")
        fake_google.genai = types.SimpleNamespace(Client=Client)
        with patch.dict(os.environ, {"GEMINI_API_KEY": "unit-test-key", "GEMINI_MODEL": "gemini-test"}), patch.dict(sys.modules, {"google": fake_google}):
            result = GeminiDocumentGenerator().generate_document("NDA", "A and B", "Preserve this term; Another clause", "Today")
        self.assertIn("Preserve this term", result)
        self.assertEqual(calls["model"], "gemini-test")
        self.assertIn("Another clause", calls["contents"])

    def test_exports_are_nonempty_and_have_valid_signatures(self):
        source = "# FREELANCE AGREEMENT\n\n## Payment\n\nPayment is due within 30 days. ‘Quotes’ and an en dash — survive.\n\n1. Confidentiality applies."
        docx = format_docx(source, "Freelance Work Contract")
        pdf = format_pdf(source, "Freelance Work Contract")
        self.assertTrue(docx.startswith(b"PK"))
        self.assertIn(b"%PDF", pdf[:10])
        self.assertGreater(len(docx), 1000)
        self.assertGreater(len(pdf), 1000)
        self.assertIn("Payment is due", "\n".join(p.text for p in Document(BytesIO(docx)).paragraphs))
        word_doc = Document(BytesIO(docx))
        self.assertTrue(any(name.startswith("word/media/") for name in ZipFile(BytesIO(docx)).namelist()))
        self.assertIn("LegalEase", word_doc.sections[0].footer.paragraphs[0].text)
        pdf_page = PdfReader(BytesIO(pdf)).pages[0]
        self.assertIn("Payment is due", pdf_page.extract_text())
        self.assertTrue(pdf_page.images)
        self.assertIn("Page 1", pdf_page.extract_text())

    def test_pdf_breaks_extremely_long_tokens_and_preserves_text(self):
        token = "LEGALCLAUSE" * 40
        pdf = format_pdf(f"# SUPPLEMENT\n\n{token}", "Supplement")
        extracted = "".join(PdfReader(BytesIO(pdf)).pages[0].extract_text().split())
        self.assertIn(token, extracted)

    def test_long_document_creates_multiple_pdf_pages(self):
        source = "# SERVICE AGREEMENT\n\n" + "\n\n".join(
            f"## Clause {number}\n\n" + ("This clause describes the obligations of both parties. " * 18)
            for number in range(1, 45)
        )
        pages = PdfReader(BytesIO(format_pdf(source, "Service Agreement"))).pages
        self.assertGreater(len(pages), 1)
        self.assertTrue(all(page.extract_text().strip() for page in pages))


if __name__ == "__main__":
    unittest.main()
