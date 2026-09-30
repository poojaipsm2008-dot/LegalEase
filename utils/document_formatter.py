"""DOCX and HTML preview formatting."""

from html import escape
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from .text_utils import parse_document

ROOT = Path(__file__).resolve().parents[1]
LOGO = ROOT / "assets" / "legalease_logo.png"


def format_docx(text: str, doc_type: str = "Legal Document") -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.9)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15

    header = section.header.paragraphs[0]
    if LOGO.is_file():
        header.add_run().add_picture(str(LOGO), width=Inches(1.35))
    else:
        header.add_run("LegalEase").bold = True
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    blocks = parse_document(text)
    first_title = next((content for kind, content in blocks if kind == "title"), doc_type)
    title_para = document.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.paragraph_format.keep_with_next = True
    title_run = title_para.add_run(first_title.upper())
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(15)

    for kind, content in blocks:
        if kind == "blank":
            continue
        if kind == "title":
            continue
        if kind == "heading2":
            para = document.add_paragraph()
            para.paragraph_format.keep_with_next = True
            run = para.add_run(content)
            run.bold = True
            run.font.size = Pt(12)
        elif kind == "list":
            para = document.add_paragraph(style="List Number")
            para.add_run(content)
        else:
            document.add_paragraph(content)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LegalEase | AI-generated draft for review")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def format_html_preview(text: str) -> str:
    """Build escaped semantic HTML for the Streamlit preview."""
    html: list[str] = []
    in_list = False
    for kind, content in parse_document(text):
        if kind != "list" and in_list:
            html.append("</ol>")
            in_list = False
        if kind == "title":
            html.append(f"<h1>{escape(content)}</h1>")
        elif kind == "heading2":
            html.append(f"<h2>{escape(content)}</h2>")
        elif kind == "paragraph":
            html.append(f"<p>{escape(content)}</p>")
        elif kind == "list":
            if not in_list:
                html.append("<ol>")
                in_list = True
            html.append(f"<li>{escape(content)}</li>")
    if in_list:
        html.append("</ol>")
    return "".join(html)
